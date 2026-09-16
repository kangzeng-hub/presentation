from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR

PATH='presentation_business_report.pptx'; prs=Presentation(PATH)
BG=RGBColor(12,23,38); PANEL=RGBColor(24,42,62); PANEL2=RGBColor(31,55,76); TEAL=RGBColor(49,205,187); ORANGE=RGBColor(255,179,82); WHITE=RGBColor(242,246,250); MUTED=RGBColor(166,184,201); GREEN=RGBColor(100,216,158); RED=RGBColor(232,112,104)
def tb(s,x,y,w,h,text,size=16,color=WHITE,bold=False,align=PP_ALIGN.LEFT):
 b=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=b.text_frame; tf.clear(); tf.word_wrap=True; tf.margin_left=Inches(.05); tf.margin_right=Inches(.05); tf.vertical_anchor=MSO_ANCHOR.MIDDLE; p=tf.paragraphs[0]; p.alignment=align; r=p.add_run(); r.text=text; r.font.name='Arial'; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=color; return b
def box(s,x,y,w,h,fill=PANEL,line=None):
 q=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h)); q.fill.solid(); q.fill.fore_color.rgb=fill; q.line.color.rgb=line or fill; return q
def line(s,x1,y1,x2,y2,c=TEAL):
 q=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2)); q.line.color.rgb=c; q.line.width=Pt(2); q.line.end_arrowhead=True
def base(title,num):
 s=prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb=BG; tb(s,.65,.28,11,.2,'PRESENTATION SYSTEM  /  ONE SKU WALKTHROUGH',10,TEAL,True); tb(s,.65,.57,12,.65,title,25,WHITE,True); tb(s,.65,7.12,9,.18,'Presentation System  |  业务汇报',9,MUTED); tb(s,12.3,7.08,.4,.2,f'{num:02d}',10,TEAL,True,PP_ALIGN.RIGHT); return s
def pic(s,path,x,y,w,h):
 try: s.shapes.add_picture(path,Inches(x),Inches(y),width=Inches(w),height=Inches(h))
 except Exception: box(s,x,y,w,h,PANEL2)
def tag(s,x,y,w,text,c): box(s,x,y,w,.3,c,c); tb(s,x,y,w,.3,text,10,BG,True,PP_ALIGN.CENTER)

# 17 full chain
s=base('一个 SKU 的完整链路：从竞品抓取一路走到 Listing、图片与视频',17)
steps=[('01','产品真相','产品资料 + 图片 + claims'),('02','竞品抓取','链接 → 标题/五点/评论/图片'),('03','竞品分析','购买驱动 / 痛点 / 差距'),('04','表达策略','优先级 + Listing/图片/视频映射'),('05','Listing','标题 + 五点 + 描述'),('06','图片','主图/橱窗图任务 → 生成 → QA'),('07','视频','场景/动作 → Prompt / VideoPlan'),('08','人工校正','确认、修改、再生成、沉淀')]
for i,(n,h,v) in enumerate(steps):
 x=.55+(i%4)*3.15; y=1.55+(i//4)*2.35; box(s,x,y,2.65,1.28,TEAL if i in [3,7] else PANEL); tb(s,x+.12,y+.14,.42,.28,n,12,BG if i in [3,7] else TEAL,True); tb(s,x+.62,y+.14,1.85,.28,h,15,BG if i in [3,7] else WHITE,True); tb(s,x+.15,y+.58,2.35,.46,v,11,BG if i in [3,7] else WHITE,False,PP_ALIGN.CENTER)
for x1,x2,y in [(3.25,3.45,2.2),(6.4,6.6,2.2),(9.55,9.75,2.2),(3.25,3.45,4.55),(6.4,6.6,4.55),(9.55,9.75,4.55)]: line(s,x1,y,x2,y,MUTED)
line(s,11.95,2.85,11.95,3.4,ORANGE); line(s,1.9,3.0,1.9,3.35,ORANGE)
tb(s,.85,6.2,11.4,.35,'同一套 ProductTruth / CompetitorInsight / PresentationStrategy 被不同内容环节引用；不是每一步重新理解产品。',16,ORANGE,True,PP_ALIGN.CENTER)

# 18 listing chain
s=base('Listing 不是单独写文案，而是把共享策略翻译成标题、五点和描述',18)
for i,(h,v) in enumerate([('竞品抓取','标题、五点、描述、价格、评分、评论'),('Insight','购买驱动、用户痛点、竞品强弱、机会差距'),('ProductTruth','真实材质、规格、结构、套装、可用 claims'),('Presentation Strategy','优先级、差异化、证据点、Listing 映射')]):
 x=.7+(i%2)*6.1; y=1.45+(i//2)*1.15; box(s,x,y,5.5,.85,PANEL); tb(s,x+.15,y+.16,1.55,.25,h,14,TEAL,True); tb(s,x+1.75,y+.16,3.5,.45,v,11,WHITE)
box(s,4.15,4.0,5.0,.7,ORANGE,ORANGE); tb(s,4.25,4.2,4.8,.26,'ListingPlan：带版本与来源引用的表达方案',15,BG,True,PP_ALIGN.CENTER)
line(s,3.25,3.2,5.0,3.92,ORANGE); line(s,9.35,3.2,8.3,3.92,ORANGE)
for i,(h,v) in enumerate([('Title','G23Studio ASTM F136 Titanium Hinged Septum Ring 3PCS…'),('Bullet Points','材质安全 / 3PCS 多款式 / 铰链闭合 / 18G·20G·8mm / 多种佩戴'),('Description','把材质、结构、规格、适用位置和套装关系组织成完整购买说明')]):
 x=.75+i*4.15; box(s,x,5.2,3.65,1.05,PANEL2); tb(s,x+.12,5.38,3.4,.25,h,14,TEAL,True,PP_ALIGN.CENTER); tb(s,x+.15,5.75,3.3,.3,v,10,WHITE,False,PP_ALIGN.CENTER)
tb(s,.85,6.65,11.5,.25,'状态：ListingPlan / mock workspace 与前端编辑界面已具备；生产级模型/provider 与正式业务审批仍需继续接入。',11,MUTED,False,PP_ALIGN.CENTER)

#19 images before/after
s=base('主图 / 橱窗图：人工接入不是最后挑图，而是改动“任务和约束”后再生成',19)
pic(s,'output/wan_baseline_v1/image_01.png',.75,1.45,2.5,2.25); pic(s,'output/gpt_image_2_from_wan/image_01.png',3.45,1.45,2.5,2.25); pic(s,'output/wan_baseline_v1/image_02.png',6.15,1.45,2.5,2.25); pic(s,'output/gpt_image_2_from_wan/image_02.png',8.85,1.45,2.5,2.25)
for x,t,c in [(.75,'首轮生成 / 主图',RED),(3.45,'人工校准后再生成 / 主图',GREEN),(6.15,'首轮生成 / 橱窗图',RED),(8.85,'人工校准后再生成 / 橱窗图',GREEN)]: tag(s,x,3.82,2.5,t,c)
box(s,.8,4.45,3.55,1.55,PANEL); tb(s,1.0,4.68,3.1,.25,'业务发现',14,ORANGE,True,PP_ALIGN.CENTER); tb(s,1.0,5.05,3.1,.68,'一眼是否看懂？\n数量/款式/证据是否正确？',14,WHITE,True,PP_ALIGN.CENTER)
box(s,4.9,4.45,3.55,1.55,PANEL); tb(s,5.1,4.68,3.1,.25,'业务输入',14,TEAL,True,PP_ALIGN.CENTER); tb(s,5.1,5.05,3.1,.68,'preserve：保留什么\nchange：改什么\ndon’t introduce：禁止什么',13,WHITE,True,PP_ALIGN.CENTER)
box(s,9.0,4.45,3.4,1.55,PANEL); tb(s,9.2,4.68,3.0,.25,'系统执行',14,TEAL,True,PP_ALIGN.CENTER); tb(s,9.2,5.05,3.0,.68,'更新 Prompt / 任务\n再生成 → QA → 再审核',14,WHITE,True,PP_ALIGN.CENTER)
line(s,4.4,5.2,4.85,5.2,ORANGE); line(s,8.5,5.2,8.95,5.2,ORANGE)
tb(s,.85,6.35,11.5,.3,'真实证据：baseline QA 为 6 张 NEEDS_REVIEW；问题码会进入 HumanEvaluation / RegenerationTask。对比图展示“再生成样例”，不宣称已通过业务验收。',11,MUTED,False,PP_ALIGN.CENTER)

#20 video
s=base('视频不是重新开始，而是复用同一套卖点与策略生成动作方案',20)
for i,(h,v) in enumerate([('共享输入','ProductTruth + Strategy + Listing/Image context'),('视频任务','目标、受众、Hook、卖点、场景、时长'),('动作脚本','展示 3PCS → 演示铰链 → 佩戴结果'),('输出','VideoPlan / Prompt / on-screen text')]):
 x=.75+i*3.1; box(s,x,1.55,2.65,1.25,TEAL if i==1 else PANEL); tb(s,x+.12,1.78,2.4,.25,h,14,BG if i==1 else TEAL,True,PP_ALIGN.CENTER); tb(s,x+.12,2.2,2.4,.38,v,11,BG if i==1 else WHITE,True,PP_ALIGN.CENTER); 
 if i<3: line(s,x+2.7,2.18,x+3.0,2.18,MUTED)
box(s,1.0,3.65,11.2,1.1,PANEL2); tb(s,1.2,3.9,10.8,.3,'当前真实 Prompt 逻辑：15 秒、三段连续动作、同一产品/模型/场景/金色表面、禁止产品变形与未验证 claims。',15,WHITE,True,PP_ALIGN.CENTER)
box(s,2.0,5.35,4.2,.7,ORANGE,ORANGE); tb(s,2.1,5.56,4.0,.25,'当前：VideoPlan / Prompt 可生成',14,BG,True,PP_ALIGN.CENTER)
box(s,7.0,5.35,4.2,.7,PANEL); tb(s,7.1,5.56,4.0,.25,'建设中：生产级视频工作区与生成验收',13,WHITE,True,PP_ALIGN.CENTER)
tb(s,.85,6.45,11.5,.25,'真实边界：`output/video/review.json` 当前 usable=false；因此 PPT 展示的是策略与 Prompt 链路，不声称最终视频已可用。',11,MUTED,False,PP_ALIGN.CENTER)

#21 close
s=base('完整链路的管理重点：校正一个上游判断，影响后面所有内容',21); box(s,.75,1.45,11.8,.8,TEAL,TEAL); tb(s,.9,1.68,11.5,.3,'例如：把“消费者最关心一套里到底有什么”确认成 P0',18,BG,True,PP_ALIGN.CENTER)
for i,(h,v) in enumerate([('竞品分析','发现竞品反复强调套装关系'),('策略','把“3PCS + 三种配置”提到优先级'),('Listing','标题/五点/描述解释套装组成'),('图片','主图/橱窗图一对一展示配置'),('视频','先展示完整 3PCS，再演示使用'),('人工校正','确认一眼看懂，必要时再生成')]):
 x=.75+(i%3)*4.15; y=2.8+(i//3)*1.55; box(s,x,y,3.65,1.05,PANEL); tb(s,x+.12,y+.18,1.0,.25,h,14,TEAL,True); tb(s,x+1.15,y+.18,2.2,.5,v,11,WHITE,True)
tb(s,.85,6.05,11.5,.42,'这就是“SKU 驱动内容生产”：一个判断被传递、执行、校正，而不是每个岗位重新开始。',18,ORANGE,True,PP_ALIGN.CENTER)

prs.save(PATH); print('appended',len(prs.slides))
