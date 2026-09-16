from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.dml import MSO_THEME_COLOR

OUT='presentation_business_report.pptx'; prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
BG=RGBColor(12,23,38); PANEL=RGBColor(24,42,62); PANEL2=RGBColor(31,55,76); TEAL=RGBColor(49,205,187); ORANGE=RGBColor(255,179,82); WHITE=RGBColor(242,246,250); MUTED=RGBColor(166,184,201); RED=RGBColor(232,112,104); GREEN=RGBColor(100,216,158)
def tb(s,x,y,w,h,text,size=16,color=WHITE,bold=False,align=PP_ALIGN.LEFT):
 b=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=b.text_frame; tf.clear(); tf.word_wrap=True; tf.margin_left=Inches(.05); tf.margin_right=Inches(.05); tf.vertical_anchor=MSO_ANCHOR.MIDDLE; p=tf.paragraphs[0]; p.alignment=align; r=p.add_run(); r.text=text; r.font.name='Arial'; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=color; return b
def box(s,x,y,w,h,fill=PANEL,line=None):
 q=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h)); q.fill.solid(); q.fill.fore_color.rgb=fill; q.line.color.rgb=line or fill; return q
def line(s,x1,y1,x2,y2,c=TEAL):
 q=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2)); q.line.color.rgb=c; q.line.width=Pt(2); q.line.end_arrowhead=True
def base(title,num):
 s=prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb=BG; tb(s,.65,.28,11,.2,'PRESENTATION SYSTEM  /  BUSINESS REVIEW',10,TEAL,True); tb(s,.65,.57,12,.65,title,26,WHITE,True); tb(s,.65,7.12,9,.18,'Presentation System  |  业务汇报',9,MUTED); tb(s,12.3,7.08,.4,.2,f'{num:02d}',10,TEAL,True,PP_ALIGN.RIGHT); return s
def tag(s,x,y,w,text,c): box(s,x,y,w,.3,c,c); tb(s,x,y,w,.3,text,10,BG,True,PP_ALIGN.CENTER)
def pic(s,path,x,y,w,h):
 try: s.shapes.add_picture(path,Inches(x),Inches(y),width=Inches(w),height=Inches(h))
 except Exception: box(s,x,y,w,h,PANEL2)

# 1
s=base('让一个 SKU 驱动完整内容生产',1); tb(s,.75,1.65,7.5,1.0,'从产品事实到 Listing、图片和视频，\n统一围绕同一套业务判断。',28,WHITE,True); tb(s,.78,3.1,6.2,.45,'系统的主角不是 AI，而是 SKU 的决策链。',18,ORANGE,True)
box(s,8.5,1.45,3.8,4.45,PANEL); tb(s,8.8,1.78,3.2,.3,'真实 SKU',17,TEAL,True,PP_ALIGN.CENTER); tb(s,8.8,2.22,3.2,.7,'G23Studio\nASTM F136 铰链鼻中隔环 3 件套',15,WHITE,True,PP_ALIGN.CENTER)
for i,t in enumerate(['产品事实','竞品输入','消费者问题','内容策略','内容资产']): box(s,8.85,3.2+i*.48,3.1,.34,BG,TEAL); tb(s,8.9,3.2+i*.48,3,.34,t,12,WHITE,True,PP_ALIGN.CENTER)

#2
s=base('为什么现有 SKU 内容生产越来越依赖人工',2); steps=['搜集资料','逐个看竞品','整理卖点','写 Listing','提图片需求','反复改图','再做视频']
for i,t in enumerate(steps): box(s,.75+i*1.72,1.85,1.35,.85,PANEL); tb(s,.8+i*1.72,2.12,1.25,.25,t,13,WHITE,True,PP_ALIGN.CENTER); 
for i in range(6): line(s,2.12+i*1.72,2.28,2.42+i*1.72,2.28,MUTED)
for i,(h,v) in enumerate([('信息散落','图片、文档、页面各自保存'),('重复理解','不同岗位重复解释产品'),('结果断裂','文案、图片、视频不一致'),('经验难复制','方法依赖个人记忆')]):
 x=.8+(i%2)*6.1; y=3.55+(i//2)*1.0; box(s,x,y,5.5,.7,BG,RGBColor(52,76,98)); tb(s,x+.18,y+.12,1.45,.25,h,14,ORANGE,True); tb(s,x+1.75,y+.12,3.45,.3,v,12,WHITE)

#3 decision chain
s=base('系统真正管理的不是内容，而是从产品判断到内容表达的决策链',3)
box(s,5.25,1.3,2.85,.6,TEAL,TEAL); tb(s,5.3,1.46,2.75,.26,'SKU',18,BG,True,PP_ALIGN.CENTER)
for x,h,v in [(1.0,'产品事实','我们到底卖什么'),(5.25,'竞品信息','市场已经强调什么'),(9.5,'用户需求','消费者为什么买/犹豫')]: box(s,x,2.45,2.85,1.0,PANEL); tb(s,x+.1,2.68,2.65,.26,h,15,TEAL,True,PP_ALIGN.CENTER); tb(s,x+.1,3.02,2.65,.24,v,12,WHITE,False,PP_ALIGN.CENTER); line(s,6.67,1.92,x+1.42,2.38,TEAL)
box(s,4.2,4.15,4.95,.75,ORANGE,ORANGE); tb(s,4.3,4.36,4.75,.3,'核心判断：这个 SKU 应该怎么卖',18,BG,True,PP_ALIGN.CENTER)
for i,t in enumerate(['Listing','图片策略','视频策略']): box(s,1.0+i*4.25,5.45,2.85,.65,PANEL2); tb(s,1.05+i*4.25,5.64,2.75,.25,t,15,WHITE,True,PP_ALIGN.CENTER); line(s,6.67,4.92,2.42+i*4.25,5.38,ORANGE)

#4
s=base('一个 SKU 进入系统以后，先发生判断，再发生生产',4); nodes=[('输入','SKU / 产品信息 / 图片 / 竞品链接 / 评论'),('判断','事实、差异、消费者问题'),('策略','优先级、表达任务、禁用项'),('执行','Listing / 图片 / 视频'),('校正','业务审核、再生成、沉淀反馈')]
for i,(h,v) in enumerate(nodes): x=.75+i*2.48; box(s,x,2.15,2.05,1.35,TEAL if i==2 else PANEL); tb(s,x+.1,2.4,1.85,.3,h,16,BG if i==2 else TEAL,True,PP_ALIGN.CENTER); tb(s,x+.1,2.82,1.85,.48,v,11,BG if i==2 else WHITE,False,PP_ALIGN.CENTER); 
for i in range(4): line(s,2.85+i*2.48,2.82,3.18+i*2.48,2.82,MUTED)
tb(s,1.0,4.45,11.2,.45,'业务人员介入点不是“最后验收”，而是可以在每个关键判断处确认、修改、否定。',19,ORANGE,True,PP_ALIGN.CENTER)

#5 product understanding
s=base('产品理解：先确定“这个 SKU 到底是什么”',5); tb(s,.8,1.45,4.0,.3,'输入什么',13,TEAL,True); tb(s,.8,1.82,4.0,1.45,'产品信息\n产品图片\n规格 / 材质 / 结构 / 限制',20,WHITE,True); line(s,4.7,2.55,5.55,2.55)
box(s,5.7,1.55,2.2,2.45,ORANGE,ORANGE); tb(s,5.82,1.85,1.95,.3,'系统判断',16,BG,True,PP_ALIGN.CENTER); tb(s,5.82,2.35,1.95,1.15,'哪些是事实\n哪些是可表达卖点\n哪些不能擅自推断',14,BG,True,PP_ALIGN.CENTER); line(s,7.95,2.55,8.7,2.55)
box(s,8.85,1.55,3.65,2.45,PANEL); tb(s,9.05,1.85,3.25,.3,'输出：统一产品认知',16,TEAL,True,PP_ALIGN.CENTER); tb(s,9.1,2.35,3.1,1.2,'ASTM F136 钛\n18K Gold PVD\n3PCS / 三种款式\n铰链、平接缝、按压卡扣',14,WHITE,True,PP_ALIGN.CENTER)
tb(s,.85,4.55,11.5,.4,'决定依据：真实产品资料 + 一方图片 + 审核通过的 claims；原则是“事实由输入决定，不由系统猜测”。',16,ORANGE,True,PP_ALIGN.CENTER)

#6 competitor
s=base('竞品研究：找到市场已经在强调什么，而不是照着竞品模仿',6); pic(s,'c1/s1.jpg',.75,1.5,2.5,2.4); pic(s,'c1/s4.jpg',3.5,1.5,2.5,2.4); pic(s,'c2/s3.jpg',6.25,1.5,2.5,2.4); box(s,9.1,1.5,3.4,2.4,PANEL); tb(s,9.3,1.77,3.0,.25,'竞品输入',15,TEAL,True,PP_ALIGN.CENTER); tb(s,9.3,2.2,3.0,1.2,'标题 / 五点\n图片构图\n1–3 星评论\n价格、评分、评价量',15,WHITE,True,PP_ALIGN.CENTER)
line(s,2.0,4.2,11.0,4.2,ORANGE); tb(s,.85,4.55,11.5,.3,'比较关系',13,ORANGE,True,PP_ALIGN.CENTER); tb(s,1.0,5.0,11.2,.55,'竞品内容 + 竞品图片 + 用户反馈 + 我方产品事实  → 竞争差异  →  内容策略依据',19,WHITE,True,PP_ALIGN.CENTER)

#7 consumer
s=base('消费者关注点：系统进一步判断用户为什么买，或者为什么犹豫',7); tb(s,.85,1.35,11.4,.35,'在当前 ImagePlan 中，消费者问题已经被写成每张图片的 decision_question 与 customer_need。',14,MUTED,False,PP_ALIGN.CENTER)
qs=[('识别','我买到的到底是哪三种款式？'),('套装','一套里面到底包含什么？'),('选择','18G / 20G / 8mm 怎么理解？'),('信任','材质和表面处理依据是什么？'),('使用','铰链卡扣到底是什么结构？'),('风格','戴上之后大概是什么风格？')]
for i,(h,v) in enumerate(qs): x=.75+(i%3)*4.15; y=2.05+(i//3)*1.6; box(s,x,y,3.65,1.12,PANEL); tb(s,x+.15,y+.18,1.0,.3,h,15,TEAL,True); tb(s,x+1.1,y+.18,2.35,.6,v,12,WHITE,True)
tb(s,.85,5.65,11.5,.45,'这些不是凭空新增的用户画像，而是项目中已落到六张图任务里的真实购买问题。',16,ORANGE,True,PP_ALIGN.CENTER)

#8 strategy
s=base('内容策略：把产品事实、竞品差异和消费者问题变成“我们应该怎么卖”',8)
for i,(h,v) in enumerate([('产品事实','能说什么'),('竞品差异','要超越什么'),('消费者问题','先回答什么'),('业务要求','必须遵守什么')]):
 x=.75+i*3.1; box(s,x,1.55,2.65,1.0,PANEL); tb(s,x+.1,1.78,2.45,.25,h,15,TEAL,True,PP_ALIGN.CENTER); tb(s,x+.1,2.15,2.45,.22,v,12,WHITE,False,PP_ALIGN.CENTER)
for i in range(3): line(s,2.1+i*3.1,2.62,2.65+i*3.1,3.1,ORANGE)
box(s,3.35,3.3,6.55,.85,ORANGE,ORANGE); tb(s,3.48,3.56,6.3,.3,'内容策略：哪些内容必须强调，哪些不要强调',17,BG,True,PP_ALIGN.CENTER)
for i,t in enumerate(['Listing：怎么写','图片：怎么展示','视频：怎么演示']): box(s,1.0+i*4.25,5.05,2.85,.7,PANEL2); tb(s,1.05+i*4.25,5.27,2.75,.25,t,14,WHITE,True,PP_ALIGN.CENTER); line(s,6.62,4.18,2.42+i*4.25,4.98,TEAL)

#9 shared decision
s=base('一个核心判断，可以同时驱动 Listing、图片和视频',9); box(s,4.25,1.35,4.8,.75,TEAL,TEAL); tb(s,4.35,1.58,4.6,.3,'核心判断：消费者最关心佩戴方式',17,BG,True,PP_ALIGN.CENTER)
for i,(t,v) in enumerate([('Listing','说明一套多种佩戴方式'),('图片','分开呈现三种真实配置'),('视频','按“展示套装 → 演示结构 → 佩戴结果”组织动作')]):
 x=.9+i*4.15; box(s,x,3.25,3.35,1.35,PANEL); tb(s,x+.15,3.52,3.05,.28,t,16,TEAL,True,PP_ALIGN.CENTER); tb(s,x+.15,3.95,3.05,.42,v,12,WHITE,True,PP_ALIGN.CENTER); line(s,6.65,2.15,x+1.68,3.18,ORANGE)
tb(s,.9,5.7,11.5,.4,'系统价值：不是分别生成三种内容，而是让同一个业务判断持续传递到三个内容环节。',17,ORANGE,True,PP_ALIGN.CENTER)

#10 image strategy
s=base('图片策略：先决定“这一张图承担什么任务”，再进入生成',10); pic(s,'p1/p1.jpg',.8,1.5,3.3,1.65); tb(s,.85,3.35,3.2,.28,'真实一方产品图：三种配置',13,MUTED,False,PP_ALIGN.CENTER)
labels=[('消费者问题','一套里到底有什么？'),('图片任务','让三种真实配置一眼可区分'),('画面要求','3PCS + 一对一标签 + 产品主导'),('业务校正','数量对不对？一眼看懂吗？')]
for i,(h,v) in enumerate(labels): x=4.65+(i%2)*3.7; y=1.55+(i//2)*1.45; box(s,x,y,3.25,1.05,PANEL if i!=3 else RGBColor(73,55,38)); tb(s,x+.13,y+.18,1.0,.23,h,13,TEAL if i!=3 else ORANGE,True); tb(s,x+1.1,y+.18,1.95,.52,v,12,WHITE,True)
tb(s,4.7,4.85,7.1,.4,'ImagePlan 已把 decision_question、customer_need、required facts、template、forbidden elements 固化。',14,ORANGE,True)

#11 case input
s=base('真实 SKU 案例：G23Studio 3 件套如何从输入走向内容任务',11); pic(s,'p1/p1.jpg',.75,1.45,4.2,2.2); box(s,5.25,1.45,7.15,2.2,PANEL); tb(s,5.5,1.7,6.65,.3,'输入事实（已实现）',15,TEAL,True); tb(s,5.5,2.15,6.55,1.1,'ASTM F136 植入级钛；18K Gold PVD；AAAAA CZ\n3PCS：classic hoop / double-layer hoop / CZ accent hoop\n18G / 20G；内径 8mm；铰链、平接缝、按压卡扣',14,WHITE,True)
for i,(h,v) in enumerate([('竞品发现','竞品反复使用分组、标签、近景和结构示意'),('消费者问题','我买到的到底是哪三种？一套包含什么？'),('策略落点','P0 先解决识别、套装内容、规格选择')]):
 y=4.15+i*.75; box(s,.85,y,11.55,.58,BG,RGBColor(53,77,99)); tb(s,1.05,y+.13,1.5,.25,h,13,ORANGE,True); tb(s,2.75,y+.13,9.35,.25,v,13,WHITE)

#12 case output
s=base('案例中的关键决策：系统明确了每一张图片为什么存在',12)
items=[('image_01','产品识别','我买到的到底是哪三种款式？','P0'),('image_02','套装内容','一套里面到底包含什么？','P0'),('image_03','规格选择','18G / 20G / 8mm 怎么理解？','P0'),('image_04','材质证据','材质和表面处理依据是什么？','P1'),('image_05','结构使用','铰链卡扣到底是什么结构？','P1'),('image_06','风格投射','戴上之后大概是什么风格？','P1')]
for i,(id,r,q,p) in enumerate(items): x=.75+(i%3)*4.15; y=1.5+(i//3)*2.1; box(s,x,y,3.65,1.5,PANEL); tb(s,x+.15,y+.16,.8,.24,id,12,TEAL,True); tb(s,x+1.05,y+.16,2.35,.24,r,15,WHITE,True); tb(s,x+.15,y+.6,3.3,.48,q,13,WHITE); tag(s,x+.15,y+1.15,.55,p,ORANGE if p=='P0' else MUTED)
tb(s,.85,6.15,11.5,.35,'策略结果不是“生成六张图”，而是把六个购买问题拆成六个可检查的内容任务。',16,ORANGE,True,PP_ALIGN.CENTER)

#13 qa
s=base('案例中的结果校正：系统先给方案，业务再判断是否真的解决问题',13); pic(s,'output/wan_baseline_v1/image_01.png',.8,1.45,2.55,2.25); pic(s,'output/wan_baseline_v1/image_02.png',3.55,1.45,2.55,2.25); box(s,6.55,1.45,5.8,2.25,PANEL); tb(s,6.8,1.72,5.3,.28,'真实 QA 结果：6 张全部 NEEDS_REVIEW',16,ORANGE,True,PP_ALIGN.CENTER); tb(s,6.85,2.2,5.2,1.05,'识别到的问题：\nDECISION_NOT_COVERED ×3\nMISSING_VISUAL_EVIDENCE ×6\nUNSUPPORTED_CLAIM ×3\nUNVERIFIED_SCALE ×5',14,WHITE,True,PP_ALIGN.CENTER)
tb(s,.85,4.35,11.5,.35,'业务校正链路',14,TEAL,True,PP_ALIGN.CENTER); flow=['系统分析','生成方案','业务审核','发现偏差','改任务/布局/表达','再生成'];
for i,t in enumerate(flow): box(s,.9+i*2.0,4.85,1.55,.7,TEAL if i==2 else PANEL); tb(s,.95+i*2.0,5.08,1.45,.23,t,12,BG if i==2 else WHITE,True,PP_ALIGN.CENTER); 
for i in range(5): line(s,2.48+i*2.0,5.2,2.83+i*2.0,5.2,MUTED)
tb(s,.9,6.15,11.4,.32,'这不是失败，而是校准机制的价值：把“看起来能生成”变成“业务问题真的被回答”。',15,ORANGE,True,PP_ALIGN.CENTER)

#14 feedback loop
s=base('每一次人工修改，都应该变成下一次生产的依据',14); loop=[('系统分析','识别偏差'),('业务审核','确认问题'),('校正变量','改卖点/任务/表达'),('再生成','验证结果'),('经验沉淀','模板/规则/案例')]
for i,(h,v) in enumerate(loop): x=.75+i*2.48; box(s,x,2.0,2.0,1.2,TEAL if i in [1,4] else PANEL); tb(s,x+.1,2.25,1.8,.25,h,14,BG if i in [1,4] else TEAL,True,PP_ALIGN.CENTER); tb(s,x+.1,2.65,1.8,.28,v,11,BG if i in [1,4] else WHITE,True,PP_ALIGN.CENTER)
for i in range(4): line(s,2.8+i*2.48,2.6,3.15+i*2.48,2.6,ORANGE)
tb(s,1.0,4.25,11.3,.55,'人工不是只在最后“挑好看的”，而是在关键决策点改变输入，影响后面的 Listing、图片和视频。',20,WHITE,True,PP_ALIGN.CENTER)
tb(s,1.3,5.55,10.7,.35,'当前已有：HumanEvaluation、问题码、再生成任务、图片版本历史、RoleSpec revision。',14,MUTED,False,PP_ALIGN.CENTER)

#15 manager controls
s=base('系统把复杂生产过程拆成了业务主管可以直接管理的决策点',15); controls=['产品卖点优先级','竞品差异判断','消费者重点关注','Listing 表达重点','主图/橱窗图任务','视频表达任务','最终结果校正']
for i,t in enumerate(controls): x=.85+(i%4)*3.05; y=1.65+(i//4)*1.35; box(s,x,y,2.65,.78,PANEL if i!=6 else ORANGE); tb(s,x+.1,y+.2,2.45,.3,t,14,BG if i==6 else WHITE,True,PP_ALIGN.CENTER)
tb(s,.9,5.35,11.5,.5,'业务主管无需关注具体生成过程，只需确认：关键判断是否正确、证据是否充分、表达是否符合业务目标。',19,ORANGE,True,PP_ALIGN.CENTER)

#16 value/status
s=base('从“人驱动内容生产”，走向“SKU 驱动内容生产”',16); tb(s,.8,1.3,11.7,.9,'当前基础已经形成；下一步是把分散能力真正串成统一工作区。',25,WHITE,True,PP_ALIGN.CENTER)
for i,(h,c,items) in enumerate([('已实现',GREEN,['产品事实/claims/品牌约束','竞品采集 MVP 与视觉分析','六图 ImagePlan、图片生成、QA、校准']),('正在建设',ORANGE,['统一 SKU 工作区与研究编排','Listing 生成/编辑/版本服务','策略一次推导、多个模块复用']),('下一阶段',TEAL,['VideoPlan 与视频工作区','端到端合同/API/前端 smoke flow','品类经验、模板、案例库沉淀'])]):
 x=.7+i*4.2; box(s,x,2.65,3.85,2.55,PANEL); tb(s,x+.2,2.93,3.45,.3,h,17,c,True,PP_ALIGN.CENTER)
 for j,it in enumerate(items): tb(s,x+.25,3.5+j*.47,3.35,.3,'• '+it,12,WHITE)
tb(s,.85,5.85,11.5,.48,'最终价值：把业务人员的经验变成可复用的规则、模板、策略与标准流程。',20,ORANGE,True,PP_ALIGN.CENTER)

prs.save(OUT); print(OUT)
