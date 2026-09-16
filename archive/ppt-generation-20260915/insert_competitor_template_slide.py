from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.shapes import MSO_CONNECTOR

PATH='presentation_business_report.pptx'; prs=Presentation(PATH)
BG=RGBColor(12,23,38); PANEL=RGBColor(24,42,62); PANEL2=RGBColor(31,55,76); TEAL=RGBColor(49,205,187); ORANGE=RGBColor(255,179,82); WHITE=RGBColor(242,246,250); MUTED=RGBColor(166,184,201); GREEN=RGBColor(100,216,158)
def tb(s,x,y,w,h,text,size=16,color=WHITE,bold=False,align=PP_ALIGN.LEFT):
 b=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=b.text_frame; tf.clear(); tf.word_wrap=True; tf.margin_left=Inches(.05); tf.margin_right=Inches(.05); tf.vertical_anchor=MSO_ANCHOR.MIDDLE; p=tf.paragraphs[0]; p.alignment=align; r=p.add_run(); r.text=text; r.font.name='Arial'; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=color; return b
def box(s,x,y,w,h,fill=PANEL,line=None):
 q=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h)); q.fill.solid(); q.fill.fore_color.rgb=fill; q.line.color.rgb=line or fill; return q
def pic(s,path,x,y,w,h):
 try: s.shapes.add_picture(path,Inches(x),Inches(y),width=Inches(w),height=Inches(h))
 except: box(s,x,y,w,h,PANEL2)

# Append then move XML into position after current slide 4.
s=prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb=BG
tb(s,.65,.28,11,.2,'PRESENTATION SYSTEM  /  COMPETITOR EVIDENCE',10,TEAL,True); tb(s,.65,.57,12,.65,'竞品比对的结果，不是“模仿谁”，而是抽象市场已经验证的表达任务',24,WHITE,True); tb(s,.65,7.12,9,.18,'Presentation System  |  业务汇报',9,MUTED); tb(s,12.3,7.08,.4,.2,'05',10,TEAL,True,PP_ALIGN.RIGHT)
pic(s,'c1/s1.jpg',.72,1.45,1.75,1.55); pic(s,'c1/s4.jpg',2.62,1.45,1.75,1.55); pic(s,'c1/s5.jpg',4.52,1.45,1.75,1.55); pic(s,'c2/s3.jpg',6.42,1.45,1.75,1.55)
tb(s,.75,3.15,7.4,.24,'竞品素材中反复出现的模式',13,ORANGE,True,PP_ALIGN.CENTER)
patterns=[('产品分组 + 白底','先解决“我买到什么”'),('近景 / 尺寸对比','解决“大小和选择”'),('结构示意 / 细节网格','解决“怎么用、凭什么信”'),('三联风格 / 场景图','解决“戴上是什么感觉”')]
for i,(h,v) in enumerate(patterns):
 x=.75+(i%2)*3.75; y=3.52+(i//2)*.78; box(s,x,y,3.45,.58,PANEL); tb(s,x+.1,y+.11,1.45,.22,h,12,TEAL,True); tb(s,x+1.58,y+.11,1.7,.24,v,10,WHITE,True)
box(s,8.65,1.45,3.8,4.35,PANEL2); tb(s,8.9,1.72,3.3,.25,'抽象成系统模板',15,TEAL,True,PP_ALIGN.CENTER)
templates=[('PRODUCT_IDENTITY_MAIN','产品识别'),('CONTENTS_SET','套装内容'),('SIZE_CHOICE_PROOF','规格选择'),('MATERIAL_EVIDENCE','材质证据'),('STRUCTURE_USE_DETAIL','结构使用'),('WEAR_STYLE_TRIPTYCH','风格投射')]
for i,(code,label) in enumerate(templates):
 y=2.15+i*.52; box(s,8.92,y,3.25,.36,BG,TEAL); tb(s,9.02,y+.03,2.05,.28,code,9,WHITE,True); tb(s,11.1,y+.03,.9,.28,label,10,ORANGE,True,PP_ALIGN.RIGHT)
tb(s,.8,6.05,11.5,.38,'模板只抽象构图、信息层级与消费者问题；产品形状、数量、claims 和品牌表达仍由我方 ProductTruth 与业务策略决定。',14,ORANGE,True,PP_ALIGN.CENTER)

# move new slide after existing slide 4 (index 3)
sldIdLst=prs.slides._sldIdLst
new_id=sldIdLst[-1]
sldIdLst.remove(new_id); sldIdLst.insert(4,new_id)
prs.save(PATH); print('inserted',len(prs.slides))
