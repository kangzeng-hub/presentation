from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
PATH='presentation_business_report.pptx'; p=Presentation(PATH); s=p.slides[4]
for sh in list(s.shapes): s.shapes._spTree.remove(sh._element)
BG=RGBColor(12,23,38); PANEL=RGBColor(24,42,62); PANEL2=RGBColor(31,55,76); TEAL=RGBColor(49,205,187); ORANGE=RGBColor(255,179,82); WHITE=RGBColor(242,246,250); MUTED=RGBColor(166,184,201); RED=RGBColor(232,112,104); GREEN=RGBColor(100,216,158)
def tb(x,y,w,h,text,size=16,color=WHITE,bold=False,align=PP_ALIGN.LEFT):
 b=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=b.text_frame; tf.clear(); tf.word_wrap=True; tf.margin_left=Inches(.05); tf.margin_right=Inches(.05); tf.vertical_anchor=MSO_ANCHOR.MIDDLE; q=tf.paragraphs[0]; q.alignment=align; r=q.add_run(); r.text=text; r.font.name='Arial'; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=color; return b
def box(x,y,w,h,fill=PANEL,line=None):
 q=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h)); q.fill.solid(); q.fill.fore_color.rgb=fill; q.line.color.rgb=line or fill; return q
def line(x1,y1,x2,y2,c=TEAL):
 q=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2)); q.line.color.rgb=c; q.line.width=Pt(2); q.line.end_arrowhead=True
def pic(path,x,y,w,h):
 try: s.shapes.add_picture(path,Inches(x),Inches(y),width=Inches(w),height=Inches(h))
 except: box(x,y,w,h,PANEL2)
s.background.fill.solid(); s.background.fill.fore_color.rgb=BG
tb(.65,.28,11,.2,'PRESENTATION SYSTEM  /  COMPETITOR EVIDENCE',10,TEAL,True); tb(.65,.57,12,.65,'竞品比对如何推导消费者关注点：从画面事实到购买问题',24,WHITE,True); tb(.65,7.12,9,.18,'Presentation System  |  业务汇报',9,MUTED); tb(12.3,7.08,.4,.2,'05',10,TEAL,True,PP_ALIGN.RIGHT)
pic('c1/s2.jpg',.72,1.42,3.2,3.2); tb(.75,4.72,3.1,.28,'具体例子：竞品 c1-s2',12,MUTED,True,PP_ALIGN.CENTER)
box(4.25,1.42,3.85,1.15,PANEL); tb(4.45,1.62,1.2,.25,'画面事实',14,TEAL,True); tb(5.7,1.62,2.15,.58,'同一鼻部近景，被切成 4 个相同构图；标签为 18g / 7mm、8mm、9mm、10mm。',11,WHITE,True)
box(4.25,2.78,3.85,1.15,PANEL); tb(4.45,2.98,1.2,.25,'可推断问题',14,ORANGE,True); tb(5.7,2.98,2.15,.58,'消费者可能在比较：不同内径戴上是什么效果？我应该选哪个尺寸？',11,WHITE,True)
box(4.25,4.14,3.85,1.15,PANEL); tb(4.45,4.34,1.2,.25,'推断依据',14,GREEN,True); tb(5.7,4.34,2.15,.58,'同一解剖位置 + 重复对比 + 明确数字标签 → 画面主动承担“选择教育”。',11,WHITE,True)
line(8.28,3.0,8.75,3.0,ORANGE)
box(8.9,1.42,3.55,3.87,PANEL2); tb(9.15,1.7,3.05,.3,'从竞品事实到我方任务',15,TEAL,True,PP_ALIGN.CENTER)
tb(9.15,2.2,3.05,.3,'竞品观察',12,MUTED,True); tb(9.15,2.53,3.05,.55,'消费者需要“尺寸 / 佩戴效果”的可视化比较。',12,WHITE,True,PP_ALIGN.CENTER)
line(10.42,3.12,10.42,3.42,ORANGE)
tb(9.15,3.48,3.05,.3,'我方 ProductTruth',12,MUTED,True); tb(9.15,3.8,3.05,.55,'18G / 20G、内径 8mm、适用位置列表；但现有主图不能证明佩戴比例。',12,WHITE,True,PP_ALIGN.CENTER)
line(10.42,4.4,10.42,4.68,ORANGE)
tb(9.15,4.73,3.05,.3,'ImagePlan 输出',12,MUTED,True); tb(9.15,5.02,3.05,.3,'SIZE_CHOICE_PROOF：先回答规格选择，再补真实佩戴证据。',11,ORANGE,True,PP_ALIGN.CENTER)
tb(.8,5.72,11.5,.48,'关键边界：竞品图片提供“消费者问题的证据”，不提供我方产品身份、尺寸或 claims；我方尺寸比例必须由 ProductTruth 与真实佩戴素材确认。',14,ORANGE,True,PP_ALIGN.CENTER)
p.save(PATH); print('rebuilt slide 5')
