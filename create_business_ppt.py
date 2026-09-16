from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR

OUT='presentation_business_report.pptx'
prs=Presentation(); prs.slide_width= Inches(13.333); prs.slide_height= Inches(7.5)
BG=RGBColor(13,24,39); PANEL=RGBColor(24,40,60); TEAL=RGBColor(46,196,182); ORANGE=RGBColor(255,174,84); WHITE=RGBColor(242,246,250); MUTED=RGBColor(166,183,199); RED=RGBColor(235,112,104); GREEN=RGBColor(94,210,150)

def tb(slide,x,y,w,h,text,size=18,color=WHITE,bold=False,align=PP_ALIGN.LEFT):
    box=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=box.text_frame; tf.clear(); tf.word_wrap=True; tf.margin_left=Inches(.04); tf.margin_right=Inches(.04); tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    p=tf.paragraphs[0]; p.alignment=align; r=p.add_run(); r.text=text; r.font.name='Arial'; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=color
    return box
def rect(slide,x,y,w,h,fill=PANEL,line=None,r=MSO_SHAPE.ROUNDED_RECTANGLE):
    s=slide.shapes.add_shape(r,Inches(x),Inches(y),Inches(w),Inches(h)); s.fill.solid(); s.fill.fore_color.rgb=fill; s.line.color.rgb=line or fill; return s
def base(title,kicker='BUSINESS REVIEW'):
    s=prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb=BG
    tb(s,.65,.32,12,.25,kicker,10,TEAL,True); tb(s,.65,.62,12,.62,title,27,WHITE,True)
    tb(s,.65,7.12,8,.2,'Presentation System  |  业务汇报',9,MUTED)
    return s
def pill(s,x,y,w,text,color=TEAL): rect(s,x,y,w,.32,color,color); tb(s,x,y,w,.32,text,10,BG,True,PP_ALIGN.CENTER)
def arrow(s,x1,y1,x2,y2,color=TEAL):
    ln=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2)); ln.line.color.rgb=color; ln.line.width=Pt(2); ln.line.end_arrowhead=True

# 1
s=base('从 SKU 到完整内容资产的一体化生产体系','PRESENTATION SYSTEM')
tb(s,.7,1.65,7.1,1.2,'让产品、竞品、Listing、图片和视频围绕同一个 SKU 协同生产。',27,WHITE,True)
tb(s,.72,3.05,5.9,.55,'从分散执行，走向统一认知、统一策略、统一复用。',17,MUTED)
rect(s,8.2,1.55,4.25,4.15,PANEL)
tb(s,8.65,1.85,3.35,.3,'一个 SKU',18,TEAL,True,PP_ALIGN.CENTER)
for i,t in enumerate(['产品事实','竞品洞察','内容策略','图片策略','视频策略']):
    rect(s,8.75,2.35+i*.55,3.15,.38,BG,TEAL); tb(s,8.8,2.35+i*.55,3.05,.38,t,13,WHITE,True,PP_ALIGN.CENTER)
tb(s,.72,5.75,11,.45,'核心判断：不是“AI 直接生成内容”，而是“先确定策略，再生成内容”。',18,ORANGE,True)

#2
s=base('一个 SKU 的内容生产，过去被拆成了多段人工协作')
steps=['找资料','找竞品','读页面','写 Listing','提图片需求','反复改图','再做视频']
for i,t in enumerate(steps):
    x=.8+i*1.72; rect(s,x,2.0,1.35,1.0,PANEL); tb(s,x+.08,2.18,1.2,.35,t,14,WHITE,True,PP_ALIGN.CENTER)
    if i<len(steps)-1: arrow(s,x+1.38,2.5,x+1.65,2.5,MUTED)
tb(s,.82,3.55,11.5,.4,'每到一个新环节，都要重新解释产品；信息分散、沟通重复、返工难以定位。',18,MUTED)
for i,(h,v) in enumerate([('信息分散','事实散落在图片、文档、页面'),('重复调研','每个岗位重复看竞品'),('表达断裂','文案、图片、视频各说各话'),('经验难复制','优秀运营的方法停留在个人')]):
    x=.85+(i%2)*6.1; y=4.45+(i//2)*.95; rect(s,x,y,5.45,.68,BG,RGBColor(50,72,94)); tb(s,x+.2,y+.12,1.55,.28,h,14,ORANGE,True); tb(s,x+1.8,y+.12,3.35,.32,v,12,WHITE)

#3
s=base('系统把分散的内容生产环节串成一条 SKU 链路')
flow=['SKU','产品事实\n与图片','竞品信息\n与低星评价','产品/竞品\n统一认知','内容策略','Listing\n图片\n视频','人工审核\n与校准','最终资产']
for i,t in enumerate(flow):
    x=.55+i*1.58; col=TEAL if i in [0,3,4] else PANEL; rect(s,x,2.0,1.28,1.15,col,col); tb(s,x+.06,2.2,1.16,.7,t,13,BG if col==TEAL else WHITE,True,PP_ALIGN.CENTER)
    if i<len(flow)-1: arrow(s,x+1.3,2.58,x+1.53,2.58,TEAL)
tb(s,.8,4.05,11.7,.35,'当前真实状态：产品事实、竞品采集 MVP、六图规划、图片生成/质检/校准已具备；Listing 工作区、统一研究编排、视频工作区仍在建设。',16,WHITE)
pill(s,.8,5.15,1.4,'已实现',GREEN); tb(s,2.35,5.15,3.5,.32,'事实、竞品、图片策略、图片 QA/校准',13,WHITE)
pill(s,6.1,5.15,1.8,'建设中',ORANGE); tb(s,8.05,5.15,4.1,.32,'统一工作区、Listing、视频计划与编排',13,WHITE)

#4
s=base('真正的变化：从“多个工具”变成“一个 SKU 驱动多个环节”')
rect(s,.8,1.7,5.25,4.4,PANEL); tb(s,1.1,2.0,4.6,.35,'过去',18,RED,True); tb(s,1.1,2.62,4.55,2.25,'Listing 团队理解一次\n设计团队重新理解一次\n视频团队再次理解一次\n工具又重新理解一次',20,WHITE,True); tb(s,1.1,5.25,4.5,.4,'结果：卖点不一致、沟通多、返工多',14,MUTED)
rect(s,7.25,1.7,5.25,4.4,RGBColor(22,52,61)); tb(s,7.55,2.0,4.6,.35,'系统之后',18,TEAL,True); tb(s,7.55,2.58,4.55,.45,'SKU → 统一产品认知',22,WHITE,True)
for i,t in enumerate(['Listing','图片','视频']): rect(s,7.7+i*1.5,3.55,1.25,.72,TEAL,TEAL); tb(s,7.72+i*1.5,3.72,1.21,.3,t,14,BG,True,PP_ALIGN.CENTER)
tb(s,7.55,5.1,4.55,.5,'核心价值：同一套事实与策略，贯穿所有内容。',15,ORANGE,True)

#5
s=base('一套核心变量，决定一个 SKU 的最终表达')
vars=[('产品事实','材质、规格、结构、套装、可用位置'),('竞品输入','竞品标题/卖点/图片模式/低星评价'),('消费者问题','识别、选择、信任、使用、风格投射'),('内容策略','先讲什么、证据强度、优先级'),('图片策略','主图任务、模板、必须展示/禁止元素'),('视频策略','演示顺序、场景、动作与一致性')]
for i,(h,v) in enumerate(vars):
    x=.75+(i%3)*4.15; y=1.65+(i//3)*2.05; rect(s,x,y,3.65,1.4,PANEL); tb(s,x+.18,y+.2,3.25,.3,h,16,TEAL if i<3 else ORANGE,True); tb(s,x+.18,y+.65,3.25,.55,v,12,WHITE)
for x in [2.6,6.75,10.9]: arrow(s,x,3.1,x,3.48,MUTED)
tb(s,.85,6.05,11.6,.4,'上游变量不是“资料仓库”，而是下游内容的决策依据：事实决定能说什么，竞品与消费者问题决定先说什么。',16,WHITE,True)

#6
s=base('从 SKU 到内容资产，当前真实工作流是“策略先行、生成跟随”')
nodes=['输入：产品目录\n+已审核声明','竞品采集\n+视觉分析','统一认知\n+消费者问题','六图 ImagePlan\n+优先级/模板/禁用项','图片生成\n+QA/人工校准','视频 Prompt\n+可生成请求','待建设：Listing\n+统一工作区']
for i,t in enumerate(nodes):
    x=.55+i*1.82; y=2.05 if i<4 else 4.1; rect(s,x,y,1.52,1.0,TEAL if i in [0,3] else PANEL,TEAL if i in [0,3] else PANEL); tb(s,x+.06,y+.14,1.4,.72,t,12,BG if i in [0,3] else WHITE,True,PP_ALIGN.CENTER)
    if i in [0,1,2]: arrow(s,x+1.55,y+.5,x+1.78,y+.5)
for x in [.55,2.37,4.19]: arrow(s,x+1.0,3.12,x+1.0,4.0,MUTED)
for x in [6.02,7.84]: arrow(s,x+1.55,4.6,x+1.78,4.6)
tb(s,.8,5.8,11.8,.45,'边界说明：视频模块目前已具备基于产品事实/卖点/消费者关注点构建 Prompt 的能力，视频工作区与完整编排尚未形成。',14,ORANGE,True)

#7
s=base('Listing、图片、视频共享同一套认知，避免三个团队各自解释产品')
rect(s,5.15,1.35,3.0,.8,TEAL,TEAL); tb(s,5.2,1.58,2.9,.28,'SKU + 统一产品认知',18,BG,True,PP_ALIGN.CENTER)
for x,t,sub in [(1.0,'Listing','卖点顺序 / 规格 / 购买理由'),(5.15,'图片','任务 / 证据 / 模板 / 禁止项'),(9.3,'视频','演示动作 / 场景 / 一致性')]:
    rect(s,x,3.35,3.0,1.2,PANEL); tb(s,x+.1,3.62,2.8,.3,t,17,TEAL,True,PP_ALIGN.CENTER); tb(s,x+.12,4.0,2.75,.3,sub,11,WHITE,False,PP_ALIGN.CENTER); arrow(s,6.65,2.2,x+1.5,3.25,TEAL)
tb(s,1.0,5.5,11.2,.5,'同一 SKU 的事实、卖点、限制条件和消费者问题被复用，而不是在每个环节重新“猜一遍”。',17,ORANGE,True,PP_ALIGN.CENTER)

#8
s=base('提效首先发生在“重复整理、重复沟通、重复试错”')
rows=[('产品信息整理','人工分散搜集','目录化、审核声明集中管理','减少重复整理'),('竞品调研','逐个打开、手工摘录','采集并标准化产品与低星评价','减少搜索与摘录'),('图片策略','从零想创意、反复解释','六图计划明确任务、证据、模板','减少沟通与试错'),('图片生成/修改','多轮尝试，问题难定位','QA + 人工校准 + 再生成','减少无效返工'),('视频准备','从零重新理解产品','复用卖点、关注点与产品事实','缩短准备时间')]
for j,h in enumerate(['环节','原方式','系统方式','提效逻辑']): tb(s,[.7,3.0,6.0,10.0][j],1.35,[2.1,2.8,3.8,2.4][j],.35,h,13,TEAL,True)
for i,row in enumerate(rows):
    y=1.85+i*.88; rect(s,.65,y,11.95,.68,BG if i%2==0 else PANEL); tb(s,.75,y+.16,2.05,.3,row[0],13,WHITE,True); tb(s,3.05,y+.16,2.75,.3,row[1],12,MUTED); tb(s,6.05,y+.16,3.75,.3,row[2],12,WHITE); tb(s,10.05,y+.16,2.35,.3,row[3],12,ORANGE,True)
tb(s,.8,6.45,11.5,.3,'说明：项目当前没有经过业务验证的节省工时数据，本页只呈现可追溯的提效逻辑，不虚构 ROI。',12,MUTED)

#9
s=base('规模化的关键，不是简单加人，而是增加系统可复用能力')
for i,(n,people,sys) in enumerate([('10 SKU','少量人工可覆盖','流程开始标准化'),('50 SKU','沟通与返工明显放大','统一输入减少重复'),('100 SKU','依赖少数熟练运营','策略与模板可复用'),('500 SKU','人力线性增长不可持续','系统承接更多重复工作')]):
    x=.75+i*3.1; rect(s,x,1.75,2.65,3.95,PANEL); tb(s,x+.18,2.05,2.25,.5,n,24,TEAL,True,PP_ALIGN.CENTER); tb(s,x+.2,2.95,2.2,.45,people,14,WHITE,True,PP_ALIGN.CENTER); tb(s,x+.2,4.15,2.2,.62,sys,14,ORANGE,True,PP_ALIGN.CENTER)
    if i<3: arrow(s,x+2.7,3.7,x+3.0,3.7,MUTED)
tb(s,.85,6.1,11.3,.4,'模型化示意，不代表已验证的人力比例；价值在于把一次性经验转成可复用流程。',14,MUTED,False,PP_ALIGN.CENTER)

#10
s=base('系统的长期价值，是把个人经验沉淀成团队生产能力')
chain=['优秀运营经验','规则与模板','SKU 工作流','案例与校准反馈','团队标准能力']
for i,t in enumerate(chain):
    x=.8+i*2.45; rect(s,x,2.3,1.9,1.15,TEAL if i in [0,4] else PANEL); tb(s,x+.1,2.58,1.7,.52,t,15,BG if i in [0,4] else WHITE,True,PP_ALIGN.CENTER)
    if i<4: arrow(s,x+1.95,2.88,x+2.35,2.88,ORANGE)
tb(s,.95,4.35,11.5,.55,'系统不是替代业务人员，而是让业务人员把时间从重复执行，转向关键判断、校准与策略优化。',20,WHITE,True,PP_ALIGN.CENTER)
tb(s,1.5,5.55,10.3,.35,'经验  →  可复制  →  可规模化  →  可持续优化',16,ORANGE,True,PP_ALIGN.CENTER)

#11
s=base('当前已具备可验证基础，但距离“全链路工作台”仍有明确建设项')
cols=[('已经可以使用',GREEN,['产品目录、规格、批准声明与品牌约束','Amazon 竞品采集 MVP + 标准化低星评价','六张图片策略计划与模板解析','图片生成、QA、人工校准、版本/再生成']),('正在建设',ORANGE,['统一项目/SKU 工作区与持久化状态','研究编排：排队、运行、完成、失败','Listing 生成/编辑/版本服务','统一策略被 Listing / 图片 / 视频共同引用']),('下一阶段',TEAL,['视频计划与工作区化生成流程','端到端合同/API 与前端工作区','更多品类经验、模板与优秀案例沉淀','以业务评审数据验证真实效率与质量'])]
for i,(h,c,items) in enumerate(cols):
    x=.65+i*4.2; rect(s,x,1.55,3.8,4.85,PANEL); tb(s,x+.22,1.85,3.35,.45,h,18,c,True); 
    for j,it in enumerate(items): tb(s,x+.25,2.65+j*.78,3.3,.56,'• '+it,13,WHITE)
tb(s,.8,6.62,11.6,.3,'状态依据：docs/system-audit.md、竞品采集/图片规划/视频模块代码与现有输出物。',11,MUTED)

#12
s=base('最终目标：从“人驱动内容生产”，走向“SKU 驱动内容生产”')
tb(s,.8,1.75,11.7,.95,'一个 SKU，拥有一套统一认知；\n一套认知，持续驱动多种内容资产。',28,WHITE,True,PP_ALIGN.CENTER)
for i,t in enumerate(['事实可信','策略一致','生产复用','人工校准','经验沉淀']):
    x=.9+i*2.45; rect(s,x,3.55,1.95,.85,TEAL if i in [0,2,4] else PANEL); tb(s,x+.08,3.8,1.8,.28,t,15,BG if i in [0,2,4] else WHITE,True,PP_ALIGN.CENTER)
tb(s,1.05,5.45,11.2,.5,'继续投入的理由：这不是一次性 AI 功能，而是在建设团队的内容生产基础设施。',20,ORANGE,True,PP_ALIGN.CENTER)
tb(s,1.8,6.25,9.8,.28,'下一步优先级：先打通流程，再提高质量，再提高自动化，最后沉淀品类能力。',14,MUTED,False,PP_ALIGN.CENTER)

prs.save(OUT)
print(OUT)
