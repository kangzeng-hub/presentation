from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
PATH='presentation_business_report.pptx'; p=Presentation(PATH); s=p.slides[8]
for sh in list(s.shapes): s.shapes._spTree.remove(sh._element)
BG=RGBColor(12,23,38); PANEL=RGBColor(24,42,62); PANEL2=RGBColor(31,55,76); TEAL=RGBColor(49,205,187); ORANGE=RGBColor(255,179,82); WHITE=RGBColor(242,246,250); MUTED=RGBColor(166,184,201); GREEN=RGBColor(100,216,158)
def tb(x,y,w,h,text,size=16,color=WHITE,bold=False,align=PP_ALIGN.LEFT):
 b=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=b.text_frame; tf.clear(); tf.word_wrap=True; tf.margin_left=Inches(.05); tf.margin_right=Inches(.05); tf.vertical_anchor=MSO_ANCHOR.MIDDLE; q=tf.paragraphs[0]; q.alignment=align; r=q.add_run(); r.text=text; r.font.name='Arial'; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=color; return b
def box(x,y,w,h,fill=PANEL,line=None):
 q=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h)); q.fill.solid(); q.fill.fore_color.rgb=fill; q.line.color.rgb=line or fill; return q
def line(x1,y1,x2,y2,c=TEAL):
 q=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2)); q.line.color.rgb=c; q.line.width=Pt(2); q.line.end_arrowhead=True
s.background.fill.solid(); s.background.fill.fore_color.rgb=BG
tb(.65,.28,11,.2,'PRESENTATION SYSTEM  /  STRATEGY DECISION',10,TEAL,True); tb(.65,.57,12,.65,'内容策略：把“发现的问题”变成“有优先级的表达方案”',25,WHITE,True); tb(.65,7.12,9,.18,'Presentation System  |  业务汇报',9,MUTED); tb(12.3,7.08,.4,.2,'09',10,TEAL,True,PP_ALIGN.RIGHT)
tb(.8,1.32,11.6,.3,'以 G23Studio 3 件套为例：不是把所有卖点都写上，而是先判断哪个购买问题最需要被解决。',15,ORANGE,True,PP_ALIGN.CENTER)
cols=[('1 事实','3PCS 三种款式\nASTM F136 钛\n铰链 / 平接缝 / 卡扣\n18G / 20G / 8mm'),('2 对比','竞品反复展示\n套装分组、尺寸近景、\n结构示意、风格三联\n我方主图暂不证明佩戴比例'),('3 问题','消费者要先确认：\n我买到什么？\n尺寸怎么选？\n材质/结构可信吗？'),('4 优先级','P0：识别 / 套装 / 规格\nP1：材质 / 结构 / 风格\nP2：生活方式后置')]
for i,(h,v) in enumerate(cols):
 x=.65+i*3.15; box(x,1.9,2.65,1.55,TEAL if i==3 else PANEL); tb(x+.12,2.12,2.4,.24,h,14,BG if i==3 else TEAL,True,PP_ALIGN.CENTER); tb(x+.15,2.52,2.35,.7,v,11,BG if i==3 else WHITE,True,PP_ALIGN.CENTER)
 if i<3: line(x+2.68,2.68,x+3.05,2.68,MUTED)
box(3.15,3.95,7.05,.72,ORANGE,ORANGE); tb(3.28,4.16,6.8,.25,'策略判断：先让用户确认“买到什么、能否选对”，再补信任与风格。',15,BG,True,PP_ALIGN.CENTER)
line(1.98,3.52,4.3,3.88,ORANGE); line(5.13,3.52,5.13,3.88,ORANGE); line(8.28,3.52,7.7,3.88,ORANGE); line(10.8,3.52,8.9,3.88,ORANGE)
tb(.8,4.95,11.5,.26,'这个判断如何落到各内容环节',13,TEAL,True,PP_ALIGN.CENTER)
rows=[('Title','产品身份 + ASTM F136 + 3PCS + 18K Gold PVD','先解决“是什么”'),('五点','材质 / 3PCS / 铰链 / 规格 / 多种佩戴','按 P0→P1 顺序回答购买问题'),('描述','把事实、结构、适用位置和限制完整展开','补充证据，不新增未审核 claims'),('主图/橱窗图','image_01–06 各自承担一个 decision_question','用模板把问题变成可检查画面'),('视频','展示 3PCS → 演示铰链 → 佩戴结果','复用同一卖点，不重新猜产品')]
for i,(h,v,r) in enumerate(rows):
 y=5.35+i*.31; box(.8,y,11.5,.27,BG if i%2==0 else PANEL); tb(.92,y+.01,1.2,.23,h,10,TEAL,True); tb(2.22,y+.01,6.4,.23,v,10,WHITE); tb(8.8,y+.01,3.2,.23,r,10,ORANGE,True)
tb(.85,6.95,11.5,.18,'业务确认点：优先级是否正确？claims 是否有证据？哪些内容必须强调/禁止表达？',11,MUTED,False,PP_ALIGN.CENTER)
p.save(PATH); print('rebuilt strategy slide')
