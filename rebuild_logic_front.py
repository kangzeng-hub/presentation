from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR

PATH='presentation_business_report.pptx'; prs=Presentation(PATH)
BG=RGBColor(12,23,38); PANEL=RGBColor(24,42,62); PANEL2=RGBColor(31,55,76); TEAL=RGBColor(49,205,187); ORANGE=RGBColor(255,179,82); WHITE=RGBColor(242,246,250); MUTED=RGBColor(166,184,201); GREEN=RGBColor(100,216,158); RED=RGBColor(232,112,104)
def clear(s):
    for sh in list(s.shapes): s.shapes._spTree.remove(sh._element)
def tb(s,x,y,w,h,text,size=16,color=WHITE,bold=False,align=PP_ALIGN.LEFT):
    b=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=b.text_frame; tf.clear(); tf.word_wrap=True; tf.margin_left=Inches(.05); tf.margin_right=Inches(.05); tf.vertical_anchor=MSO_ANCHOR.MIDDLE; p=tf.paragraphs[0]; p.alignment=align; r=p.add_run(); r.text=text; r.font.name='Arial'; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=color; return b
def box(s,x,y,w,h,fill=PANEL,line=None):
    q=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h)); q.fill.solid(); q.fill.fore_color.rgb=fill; q.line.color.rgb=line or fill; return q
def line(s,x1,y1,x2,y2,c=TEAL,arrow=True):
    q=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2)); q.line.color.rgb=c; q.line.width=Pt(2); q.line.end_arrowhead=arrow
def header(s,title,num,kicker='PRESENTATION SYSTEM  /  BUSINESS LOGIC'):
    s.background.fill.solid(); s.background.fill.fore_color.rgb=BG; tb(s,.65,.28,11,.2,kicker,10,TEAL,True); tb(s,.65,.57,12,.65,title,25,WHITE,True); tb(s,.65,7.12,9,.18,'Presentation System  |  业务汇报',9,MUTED); tb(s,12.3,7.08,.4,.2,f'{num:02d}',10,TEAL,True,PP_ALIGN.RIGHT)

# P1: cover as logic statement
s=prs.slides[0]; clear(s); header(s,'一个 SKU，先被理解，再被生产；产出之后，再被校正',1)
tb(s,.8,1.42,11.7,.52,'Presentation System 的核心不是“生成更多内容”，而是让一个业务判断贯穿一整条内容链。',19,ORANGE,True,PP_ALIGN.CENTER)
box(s,.85,2.55,5.15,2.55,PANEL); tb(s,1.1,2.82,4.65,.3,'0 → 1  建立表达',17,TEAL,True,PP_ALIGN.CENTER); tb(s,1.2,3.35,4.45,1.1,'SKU → 产品事实 → 竞品差异\n→ 消费者问题 → 内容策略\n→ Listing / 图片 / 视频',19,WHITE,True,PP_ALIGN.CENTER)
box(s,7.3,2.55,5.15,2.55,PANEL2); tb(s,7.55,2.82,4.65,.3,'1 → 0  回到业务',17,ORANGE,True,PP_ALIGN.CENTER); tb(s,7.65,3.35,4.45,1.1,'内容结果 → 人工审核 → 发现偏差\n→ 修正卖点 / 任务 / 约束\n→ 再生成 → 经验沉淀',19,WHITE,True,PP_ALIGN.CENTER)
line(s,6.05,3.8,7.22,3.8,ORANGE); tb(s,5.85,4.08,1.65,.26,'双向闭环',12,ORANGE,True,PP_ALIGN.CENTER)
tb(s,.85,5.75,11.5,.35,'前半段讲清楚“怎么从 0 变成内容”，后半段讲清楚“怎么从结果回到业务判断”。',17,WHITE,True,PP_ALIGN.CENTER)

# P2: 0->1 as one chain
s=prs.slides[1]; clear(s); header(s,'0 → 1：把分散的业务输入，收敛成一个“应该怎么卖”的判断',2)
nodes=[('SKU','一个具体产品'),('输入','产品资料 / 产品图\n竞品链接 / 评论'),('比较','我方事实 × 市场表达\n× 消费者问题'),('判断','优先级 / 差异点\n证据与禁用项'),('表达','标题五点描述\n主图橱窗图视频')]
for i,(h,v) in enumerate(nodes):
 x=.55+i*2.58; fill=TEAL if i==3 else PANEL; box(s,x,2.25,2.15,1.35,fill); tb(s,x+.1,2.48,1.95,.27,h,16,BG if i==3 else TEAL,True,PP_ALIGN.CENTER); tb(s,x+.12,2.88,1.9,.47,v,11,BG if i==3 else WHITE,True,PP_ALIGN.CENTER)
 if i<4: line(s,x+2.18,2.92,x+2.5,2.92,MUTED)
tb(s,.85,4.45,11.5,.4,'关键不是“输入越多越好”，而是让输入在进入下一步前完成业务判断。',18,ORANGE,True,PP_ALIGN.CENTER)
for i,(h,v) in enumerate([('产品事实','决定能说什么'),('竞品差异','决定要超越什么'),('消费者问题','决定先回答什么')]):
 x=.95+i*4.15; box(s,x,5.3,3.55,.7,PANEL2); tb(s,x+.12,5.5,1.1,.25,h,13,TEAL,True); tb(s,x+1.3,5.5,2.05,.25,v,12,WHITE,True)

# P3: decisions not modules
s=prs.slides[2]; clear(s); header(s,'0 → 1 的核心：每向下一步走一次，就做出一个业务决策',3)
box(s,5.25,1.25,2.85,.65,TEAL,TEAL); tb(s,5.35,1.46,2.65,.25,'SKU',18,BG,True,PP_ALIGN.CENTER)
items=[('产品真实吗？','事实 / claims / 证据'),('市场怎么卖？','竞品标题、图片、评论'),('用户在意什么？','识别、选择、信任、使用'),('我们先强调什么？','优先级与表达任务'),('内容怎么证明？','标题 / 图片 / 动作')]
for i,(h,v) in enumerate(items):
 x=.65+i*2.52; y=2.55 if i<3 else 4.45; box(s,x,y,2.18,1.02,PANEL if i!=3 else ORANGE); tb(s,x+.1,y+.18,1.98,.25,h,13,BG if i==3 else TEAL,True,PP_ALIGN.CENTER); tb(s,x+.1,y+.54,1.98,.28,v,10,BG if i==3 else WHITE,True,PP_ALIGN.CENTER)
 if i==0: line(s,6.68,1.92,1.74,2.48,TEAL)
 if i==1: line(s,6.68,1.92,4.26,2.48,TEAL)
 if i==2: line(s,6.68,1.92,6.78,2.48,TEAL)
 if i==3: line(s,1.74,3.65,1.74,4.38,ORANGE)
 if i==4: line(s,4.26,3.65,4.26,4.38,ORANGE)
tb(s,7.0,4.62,5.1,.72,'每一个框都是可被业务确认、修改或否定的决策点。',17,WHITE,True,PP_ALIGN.CENTER)

# P4: 1->0 feedback
s=prs.slides[3]; clear(s); header(s,'1 → 0：内容结果不是终点，而是反向检验业务判断',4)
chain=[('内容产出','标题 / 五点 / 描述\n主图 / 橱窗图 / 视频'),('业务审核','是否准确？\n是否一眼看懂？'),('发现偏差','事实错 / 任务错\n证据不足 / 表达不清'),('修正变量','改卖点优先级\n改图片任务 / Prompt 约束'),('再生成','新版本 → QA → 再审核')]
for i,(h,v) in enumerate(chain):
 x=.65+i*2.55; fill=ORANGE if i==2 else PANEL; box(s,x,2.05,2.15,1.35,fill); tb(s,x+.1,2.28,1.95,.25,h,14,BG if i==2 else TEAL,True,PP_ALIGN.CENTER); tb(s,x+.12,2.68,1.9,.48,v,11,BG if i==2 else WHITE,True,PP_ALIGN.CENTER); 
 if i<4: line(s,x+2.18,2.73,x+2.48,2.73,MUTED)
tb(s,.85,4.45,11.5,.45,'人工接入的位置：不是只选“好看的结果”，而是定位“哪一个业务判断需要改”。',18,ORANGE,True,PP_ALIGN.CENTER)
for i,(h,v) in enumerate([('保留 preserve','哪些事实、结构、布局不能动'),('修改 change','哪一个任务、卖点或表达要调整'),('禁止 don’t introduce','哪些新 claims、数量、形态不能出现')]):
 x=.85+i*4.15; box(s,x,5.35,3.55,.72,PANEL2); tb(s,x+.12,5.55,1.45,.25,h,12,TEAL,True); tb(s,x+1.62,5.55,1.75,.25,v,10,WHITE,True)

# P5: combined loop
s=prs.slides[4]; clear(s); header(s,'一个完整闭环：从 0 到 1，再从 1 回到 0',5)
box(s,5.25,1.15,2.85,.6,TEAL,TEAL); tb(s,5.35,1.34,2.65,.24,'SKU / 业务目标',16,BG,True,PP_ALIGN.CENTER)
top=[('事实','产品能证明什么'),('差异','市场还缺什么'),('策略','我们先强调什么'),('资产','内容如何表达')]
for i,(h,v) in enumerate(top):
 x=.85+i*3.05; box(s,x,2.3,2.45,1.05,TEAL if i==2 else PANEL); tb(s,x+.12,2.52,2.2,.24,h,14,BG if i==2 else TEAL,True,PP_ALIGN.CENTER); tb(s,x+.12,2.88,2.2,.25,v,11,BG if i==2 else WHITE,True,PP_ALIGN.CENTER)
 if i<3: line(s,x+2.48,2.82,x+2.92,2.82,MUTED)
line(s,6.68,1.8,2.05,2.22,TEAL); line(s,6.68,1.8,11.2,2.22,TEAL)
box(s,4.55,4.35,4.25,.72,ORANGE,ORANGE); tb(s,4.65,4.56,4.05,.25,'业务审核：是否真的解决购买问题？',14,BG,True,PP_ALIGN.CENTER)
line(s,2.05,3.48,5.8,4.28,ORANGE); line(s,5.8,5.1,2.05,5.82,ORANGE); line(s,7.8,5.1,11.2,5.82,ORANGE)
tb(s,.8,5.68,3.0,.3,'0 → 1：建立内容',14,TEAL,True,PP_ALIGN.CENTER); tb(s,9.55,5.68,3.0,.3,'1 → 0：修正判断',14,ORANGE,True,PP_ALIGN.CENTER)
tb(s,.85,6.45,11.5,.25,'这条双向链条，比“竞品模块 / Listing 模块 / 图片模块 / 视频模块”更接近系统的真实业务价值。',13,MUTED,False,PP_ALIGN.CENTER)

prs.save(PATH); print('rebuilt front',len(prs.slides))
