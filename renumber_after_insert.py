from pptx import Presentation
PATH='presentation_business_report.pptx'; p=Presentation(PATH)
for idx,s in enumerate(p.slides):
    if idx < 4: continue
    # slide footer number is a short exact two-digit text box
    target=f'{idx+1:02d}'
    for sh in s.shapes:
        if hasattr(sh,'text') and sh.text.strip().isdigit() and len(sh.text.strip())==2:
            sh.text_frame.paragraphs[0].runs[0].text=target
            break
p.save(PATH); print('renumbered',len(p.slides))
